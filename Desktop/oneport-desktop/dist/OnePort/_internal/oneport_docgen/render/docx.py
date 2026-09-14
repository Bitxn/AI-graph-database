# SPDX-License-Identifier: Apache-2.0
# Copyright 2024 OnePort Contributors
"""
Render a Document to an editable Microsoft Word (.docx) file.

This is the enterprise wedge: compliance and management teams live in Word /
SharePoint, and almost every dev-doc tool stops at Markdown/HTML. The output
carries the same apparatus as the PDF — cover, document-control table, revision
history, classification in the page header, live "Page X of Y" fields in the
footer, numbered sections and a sign-off table — but every word is editable.
"""
from __future__ import annotations

from pathlib import Path

from oneport_docgen.document import Block, Document


def render_docx(doc: Document, out_path: Path, brand=None) -> Path:
    from docx import Document as Docx
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    from oneport_docgen.config import Brand
    b = brand or Brand()
    m = doc.meta
    primary = _rgb(b.primary)
    accent = _rgb(b.accent)
    accent2 = _rgb(b.accent2)

    d = Docx()

    # base style
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # ---- header (classification) & footer (page X of Y) ----------------------
    sect = d.sections[0]
    sect.top_margin = Cm(2.2)
    sect.bottom_margin = Cm(1.8)
    hdr = sect.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrun = hdr.add_run(f"{(b.company + '  ·  ') if b.company else ''}"
                       f"{m.classification.upper()}  ·  {m.doc_type} v{m.version}")
    hrun.font.size = Pt(7.5)
    hrun.font.color.rgb = _rgb("#64748B")

    ftr = sect.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(ftr, f"{b.footer or m.title[:70]}    |    Page ", 7.5, "#64748B")
    _field(ftr, "PAGE", qn, OxmlElement)
    _run(ftr, " of ", 7.5, "#64748B")
    _field(ftr, "NUMPAGES", qn, OxmlElement)

    # ---- cover ---------------------------------------------------------------
    d.add_paragraph()
    d.add_paragraph()
    if b.company:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, b.company.upper(), 13, b.accent, bold=True, spacing=2)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, m.project, 30, b.primary, bold=True)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, m.doc_type, 16, b.accent2, bold=True)
    d.add_paragraph()
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, m.classification.upper(), 11, "#FFFFFF", bold=True)
    _shade_paragraph(p, b.accent2, qn, OxmlElement)
    d.add_paragraph()
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, f"Version {m.version}   ·   {m.status}   ·   {m.doc_date}", 11, "#64748B")
    d.add_page_break()

    # ---- document control ----------------------------------------------------
    _heading(d, "Document Control", accent)
    _kv_table(d, m.control_rows(), primary, qn, OxmlElement)
    d.add_paragraph()
    _heading(d, "Revision History", accent)
    _grid_table(d, [["Version", "Date", "Author", "Notes"]] +
                [[r.version, r.date, r.author, r.notes] for r in m.revisions],
                accent, qn, OxmlElement)
    d.add_page_break()

    # ---- at a glance + contents ----------------------------------------------
    if doc.at_a_glance:
        _heading(d, "At a Glance", accent)
        _kv_table(d, list(doc.at_a_glance.items()), primary, qn, OxmlElement)
        d.add_paragraph()
    _heading(d, "Table of Contents", accent)
    for num, heading in doc.toc():
        p = d.add_paragraph()
        _run(p, f"{num}.  ", 10.5, b.primary, bold=True)
        _run(p, heading, 10.5, "#1E293B")
    d.add_page_break()

    # ---- sections ------------------------------------------------------------
    for s in doc.sections:
        _heading(d, f"{s.number}. {s.heading}", accent)
        for blk in s.blocks:
            _render_block(d, blk, accent, qn, OxmlElement)

    # ---- sign-off ------------------------------------------------------------
    d.add_page_break()
    _heading(d, "Approval & Sign-off", accent)
    d.add_paragraph("This document requires the following approvals before it is considered final.")
    rows = [["Role", "Name", "Signature", "Date"]]
    for role, people in (("Author", m.authors), ("Reviewer", m.reviewers), ("Approver", m.approvers)):
        for person in (people or ["—"]):
            rows.append([role, person, "", ""])
    _grid_table(d, rows, primary, qn, OxmlElement)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------- #
# helpers                                                                        #
# --------------------------------------------------------------------------- #
def _rgb(hex_str: str):
    from docx.shared import RGBColor
    return RGBColor.from_string(hex_str.lstrip("#").upper())


def _run(paragraph, text, size, color, bold=False, spacing=0):
    from docx.shared import Pt
    r = paragraph.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = _rgb(color) if isinstance(color, str) else color
    return r


def _heading(d, text, color):
    from docx.shared import Pt
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = color if not isinstance(color, str) else _rgb(color)
    p.space_after = Pt(4)
    return p


def _field(paragraph, code, qn, OxmlElement):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = f" {code} "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)


def _shade_cell(cell, hex_color, qn, OxmlElement):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), str(hex_color).lstrip("#").upper())
    tcPr.append(shd)


def _shade_paragraph(paragraph, hex_color, qn, OxmlElement):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), str(hex_color).lstrip("#").upper())
    pPr.append(shd)


def _kv_table(d, rows, label_color, qn, OxmlElement):
    from docx.shared import Pt
    t = d.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        c0, c1 = t.rows[i].cells
        c0.text = ""; c1.text = ""
        r0 = c0.paragraphs[0].add_run(str(k)); r0.bold = True; r0.font.size = Pt(9.5)
        _shade_cell(c0, "#F1F5F9", qn, OxmlElement)
        r1 = c1.paragraphs[0].add_run(str(v)); r1.font.size = Pt(9.5)
    _set_col_widths(t, [4.6, 12.0])
    return t


def _grid_table(d, rows, head_color, qn, OxmlElement):
    from docx.shared import Pt
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for j, cell in enumerate(t.rows[0].cells):
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(rows[0][j])); run.bold = True
        run.font.size = Pt(9); run.font.color.rgb = _rgb("#FFFFFF")
        _shade_cell(cell, head_color, qn, OxmlElement)
    for i in range(1, len(rows)):
        for j, cell in enumerate(t.rows[i].cells):
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(rows[i][j])); run.font.size = Pt(8.8)
            if i % 2 == 0:
                _shade_cell(cell, "#F1F5F9", qn, OxmlElement)
    return t


def _set_col_widths(table, widths_cm):
    from docx.shared import Cm
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def _add_inline(paragraph, text, size=None):
    """Render **bold** spans as real bold runs (the shared inline convention)."""
    import re
    from docx.shared import Pt
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):
        if not part:
            continue
        r = paragraph.add_run(part)
        if size is not None:
            r.font.size = Pt(size)
        if i % 2 == 1:
            r.bold = True


def _render_block(d, blk: Block, accent, qn, OxmlElement):
    from docx.shared import Pt
    if blk.kind == "para":
        _add_inline(d.add_paragraph(), blk.text)
    elif blk.kind == "subheading":
        p = d.add_paragraph()
        r = p.add_run(blk.text); r.bold = True; r.font.size = Pt(11.5)
        r.font.color.rgb = accent if not isinstance(accent, str) else _rgb(accent)
    elif blk.kind == "code":
        p = d.add_paragraph()
        r = p.add_run(blk.text); r.font.name = "Consolas"; r.font.size = Pt(9)
    elif blk.kind == "bullets":
        for it in blk.items:
            _add_inline(d.add_paragraph(style="List Bullet"), str(it))
    elif blk.kind == "table" and blk.rows:
        _grid_table(d, blk.rows, accent if isinstance(accent, str) else "#4F46E5",
                    qn, OxmlElement)
